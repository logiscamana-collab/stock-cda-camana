# -*- coding: utf-8 -*-
"""Stock CDA Camana - actualizador automatico.
Lee los productos desde Ingreso Inv. Sap y publica data.json.
Los filtros iniciales se toman de la hoja Stock Ventas (tabla dinamica).
"""
from pathlib import Path
import json, subprocess, sys, time
from datetime import datetime
try:
    from openpyxl import load_workbook
except ImportError:
    print("Falta openpyxl. Ejecuta instalar.bat")
    sys.exit(1)
ROOT=Path(__file__).resolve().parent
CONFIG=json.loads((ROOT/'config.json').read_text(encoding='utf-8'))
EXCEL=ROOT/CONFIG['excel_filename']; DATA=ROOT/'data.json'
SOURCE_SHEET=CONFIG.get('sheet_name','Ingreso Inv. Sap')
PIVOT_SHEET=CONFIG.get('pivot_sheet_name','Stock Ventas')
POLL=int(CONFIG.get('poll_seconds',5)); COMMIT=CONFIG.get('commit_message','Actualizar Stock Ventas')
REPO_URL=f"https://github.com/{CONFIG['github_username']}/{CONFIG['github_repository']}.git"

def norm(v):
    if v is None: return ''
    return ' '.join(str(v).strip().upper().split())

def clean_number(v):
    if v is None or v=='': return None
    if isinstance(v,(int,float)):
        return int(v) if float(v).is_integer() else v
    try:
        n=float(str(v).strip().replace(',','.'))
        return int(n) if n.is_integer() else n
    except: return v

def wait_for_excel():
    last=None; stable=0
    for _ in range(12):
        try:
            size=EXCEL.stat().st_size
            if size==last:
                stable+=1
                if stable>=2: return True
            else: stable=0
            last=size
        except: pass
        time.sleep(1)
    return EXCEL.exists()

def find_header(ws, required):
    for i,row in enumerate(ws.iter_rows(values_only=True)):
        vals=[norm(x) for x in row]
        if all(x in vals for x in required): return i, vals
    return None,None

def extract_source(wb):
    if SOURCE_SHEET not in wb.sheetnames: raise ValueError(f"No existe la hoja '{SOURCE_SHEET}'.")
    ws=wb[SOURCE_SHEET]
    hidx,vals=find_header(ws,['BASIS','SAP','DESCRIPCIÓN','STOCK','FAMILIA','STATUS'])
    if hidx is None: raise ValueError('No encontré los encabezados BASIS/SAP/DESCRIPCIÓN/STOCK/Familia/Status.')
    def col(*names):
        for name in names:
            if name in vals: return vals.index(name)
        return None
    cb=col('BASIS'); cs=col('SAP'); cd=col('DESCRIPCIÓN'); cst=col('STOCK'); cf=col('FAMILIA'); cstatus=col('STATUS')
    records=[]; seen=set()
    for r in ws.iter_rows(min_row=hidx+2,values_only=True):
        desc=r[cd] if cd<len(r) else None; fam=r[cf] if cf<len(r) else None; status=r[cstatus] if cstatus<len(r) else None
        if desc in (None,'',0) or fam in (None,'') or status in (None,''): continue
        if norm(desc) in ('TOTAL','TOTAL GENERAL'): continue
        rec={'BASIS':clean_number(r[cb]),'SAP':clean_number(r[cs]),'DESCRIPCIÓN':str(desc).strip(),'STOCK PT':clean_number(r[cst]),'Familia':str(fam).strip(),'Status':str(status).strip()}
        key=(rec['BASIS'],rec['SAP'],rec['DESCRIPCIÓN'],rec['STOCK PT'],rec['Familia'],rec['Status'])
        if key not in seen: seen.add(key); records.append(rec)
    return records

def extract_pivot_filters(wb):
    filters={}
    if PIVOT_SHEET not in wb.sheetnames: return filters
    ws=wb[PIVOT_SHEET]
    for r in ws.iter_rows(min_row=1,max_row=min(ws.max_row,20),values_only=True):
        if len(r)>=2 and r[0] not in (None,'') and r[1] not in (None,''):
            k=str(r[0]).strip(); v=str(r[1]).strip()
            if k in ('Status','Familia'): filters[k]=v
    return filters

def extract_stock():
    if not EXCEL.exists(): raise FileNotFoundError(f'No se encontró: {EXCEL.name}')
    wb=load_workbook(EXCEL,data_only=True,read_only=True)
    try:
        rows=extract_source(wb); filters=extract_pivot_filters(wb)
    finally: wb.close()
    families=sorted({r['Familia'] for r in rows},key=str.casefold)
    statuses=sorted({r['Status'] for r in rows},key=str.casefold)
    return {'sheet':SOURCE_SHEET,'pivot_sheet':PIVOT_SHEET,'updated':datetime.now().isoformat(timespec='seconds'),'filters':filters,'options':{'Familia':families,'Status':statuses},'rows':rows}

def git(*args,check=True):
    p=subprocess.run(['git',*args],cwd=ROOT,text=True,capture_output=True,encoding='utf-8',errors='replace')
    if check and p.returncode!=0: raise RuntimeError((p.stdout+'\n'+p.stderr).strip())
    return p

def publish():
    if git('rev-parse','--show-toplevel',check=False).returncode!=0: raise RuntimeError('Esta carpeta no es un repositorio Git.')
    rem=git('remote','get-url','origin',check=False)
    if rem.returncode!=0: git('remote','add','origin',REPO_URL)
    git('add','data.json')
    if git('diff','--cached','--quiet',check=False).returncode==0:
        print('Sin cambios en data.json.'); return False
    git('commit','-m',COMMIT)
    p=git('push','-u','origin','main',check=False)
    if p.returncode!=0: raise RuntimeError(p.stdout+'\n'+p.stderr)
    print('Publicado en GitHub correctamente.')
    return True

def update_once():
    wait_for_excel(); payload=extract_stock(); new=json.dumps(payload,ensure_ascii=False,indent=2)+'\n'; old=DATA.read_text(encoding='utf-8') if DATA.exists() else ''
    if new==old: print(f'[{datetime.now():%H:%M:%S}] Sin cambios.'); return False
    DATA.write_text(new,encoding='utf-8')
    print(f"[{datetime.now():%H:%M:%S}] Datos actualizados: {len(payload['rows'])} productos. Filtros iniciales: {payload['filters']}")
    return publish()

def main():
    print('='*65); print(' STOCK CDA CAMANA - ACTUALIZADOR AUTOMÁTICO'); print('='*65)
    print('Excel:',EXCEL); print('Fuente:',SOURCE_SHEET); print('Filtros:',PIVOT_SHEET); print('GitHub:',REPO_URL); print(f'Revisión cada {POLL} segundos.'); print('No cierres esta ventana mientras quieras monitorear el Excel.\n')
    last=None
    try:
        if EXCEL.exists(): update_once(); last=EXCEL.stat().st_mtime_ns
    except Exception as e: print('ERROR inicial:',e)
    while True:
        try:
            if EXCEL.exists():
                m=EXCEL.stat().st_mtime_ns
                if last is None or m!=last:
                    time.sleep(2); update_once(); last=EXCEL.stat().st_mtime_ns
            else: print(f'[{datetime.now():%H:%M:%S}] Esperando Excel: {EXCEL.name}')
        except KeyboardInterrupt: print('\nPrograma detenido.'); break
        except Exception as e: print(f'[{datetime.now():%H:%M:%S}] ERROR:',e); time.sleep(5)
        time.sleep(POLL)
if __name__=='__main__': main()
